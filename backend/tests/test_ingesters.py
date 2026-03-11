from __future__ import annotations

import asyncio
from io import BytesIO
from pathlib import Path

import fitz  # type: ignore[import]
import pytest
from docx import Document as DocxDocument  # type: ignore[import]
from openpyxl import Workbook  # type: ignore[import]

from app.services.ingestion.docx_ingester import DocxIngester
from app.services.ingestion.ods_ingester import OdsIngester
from app.services.ingestion.pdf_ingester import PdfIngester
from app.services.ingestion.xlsx_ingester import XlsxIngester
from app.utils.crypto import encrypt


def _create_sample_pdf_bytes(text: str) -> bytes:
    """Create a simple one-page PDF containing the given text."""

    doc = fitz.open()
    try:
        page = doc.new_page()
        page.insert_text((72, 72), text)
        return doc.tobytes()
    finally:
        doc.close()


def _create_sample_docx_bytes(text: str) -> bytes:
    """Create a simple DOCX document containing the given text."""

    document = DocxDocument()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _create_sample_xlsx_bytes(values: list[list[str]]) -> bytes:
    """Create a simple XLSX workbook from a 2D list of strings."""

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in values:
        ws.append(row)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _write_encrypted_file(tmp_path: Path, filename: str, plaintext: bytes) -> Path:
    """Encrypt ``plaintext`` and write it to ``tmp_path/filename``."""

    encrypted = encrypt(plaintext)
    file_path = tmp_path / filename
    file_path.write_bytes(encrypted)
    return file_path


@pytest.mark.asyncio
async def test_pdf_ingester_extracts_text(tmp_path: Path) -> None:
    """PdfIngester should correctly extract text from a simple PDF."""

    sample_text = "Hello from PDF ingester"
    pdf_bytes = _create_sample_pdf_bytes(sample_text)
    encrypted_path = _write_encrypted_file(tmp_path, "sample.pdf.enc", pdf_bytes)

    ingester = PdfIngester()
    result = await ingester.ingest(
        encrypted_path,
        mime_type="application/pdf",
        file_format="pdf",
    )

    assert sample_text in result.text
    assert result.metadata.get("page_count") == 1
    assert result.metadata.get("mime_type") == "application/pdf"
    assert result.metadata.get("file_format") == "pdf"


@pytest.mark.asyncio
async def test_docx_ingester_extracts_text(tmp_path: Path) -> None:
    """DocxIngester should correctly extract text from a simple DOCX."""

    sample_text = "Hello from DOCX ingester"
    docx_bytes = _create_sample_docx_bytes(sample_text)
    encrypted_path = _write_encrypted_file(tmp_path, "sample.docx.enc", docx_bytes)

    ingester = DocxIngester()
    result = await ingester.ingest(
        encrypted_path,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_format="docx",
    )

    assert sample_text in result.text
    assert result.metadata.get("paragraph_count") == 1
    assert result.metadata.get("file_format") == "docx"


@pytest.mark.asyncio
async def test_xlsx_ingester_extracts_text(tmp_path: Path) -> None:
    """XlsxIngester should correctly extract text from a simple XLSX."""

    values = [["Name", "Value"], ["foo", "42"]]
    sample_value = "foo"
    xlsx_bytes = _create_sample_xlsx_bytes(values)
    encrypted_path = _write_encrypted_file(tmp_path, "sample.xlsx.enc", xlsx_bytes)

    ingester = XlsxIngester()
    result = await ingester.ingest(
        encrypted_path,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file_format="xlsx",
    )

    assert sample_value in result.text
    assert result.metadata.get("sheet_count") == 1
    sheets_meta = result.metadata.get("sheets") or []
    assert len(sheets_meta) == 1
    assert sheets_meta[0].get("sheet_name") == "Sheet1"


def _create_sample_ods_bytes(values: list[list[str]]) -> bytes:
    """Create a simple ODS workbook from a 2D list of strings using pandas."""

    import pandas as pd

    df = pd.DataFrame(values)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="odf") as writer:
        df.to_excel(writer, sheet_name="Sheet1", index=False, header=False)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_ods_ingester_extracts_text(tmp_path: Path) -> None:
    """OdsIngester should correctly extract text from a simple ODS."""

    values = [["Name", "Value"], ["bar", "99"]]
    sample_value = "bar"
    ods_bytes = _create_sample_ods_bytes(values)
    encrypted_path = _write_encrypted_file(tmp_path, "sample.ods.enc", ods_bytes)

    ingester = OdsIngester()
    result = await ingester.ingest(
        encrypted_path,
        mime_type="application/vnd.oasis.opendocument.spreadsheet",
        file_format="ods",
    )

    assert sample_value in result.text
    assert result.metadata.get("sheet_count") == 1
    sheets_meta = result.metadata.get("sheets") or []
    assert len(sheets_meta) == 1
    assert sheets_meta[0].get("sheet_name") == "Sheet1"

