"""Regression tests: ingesting a document never writes its decrypted bytes to disk.

Each test writes the encrypted fixture first, then installs a guard that
records every attempt to create or open a file for writing (``open``,
``os.open`` with a write flag and the :mod:`tempfile` factories) while the
ingester runs. Attempts are recorded rather than raised because some
ingesters deliberately degrade on optional steps, which would hide a raised
guard error.
"""

from __future__ import annotations

import asyncio
import os
import shutil
import subprocess
import wave
from io import BytesIO
from pathlib import Path
from typing import Any, List

import fitz
import numpy as np
import pytest
from disk_write_guard import install_write_guard
from docx import Document as DocxDocument
from openpyxl import Workbook
from pdfplumber.utils.exceptions import PdfminerException
from PIL import Image

from septum_api.services.ingestion import audio_ingester as audio_module
from septum_api.services.ingestion import image_ingester as image_module
from septum_api.services.ingestion import pdf_ingester as pdf_module
from septum_api.services.ingestion.audio_ingester import AudioDecodeError, AudioIngester
from septum_api.services.ingestion.docx_ingester import DocxIngester
from septum_api.services.ingestion.image_ingester import ImageIngester
from septum_api.services.ingestion.ods_ingester import OdsIngester
from septum_api.services.ingestion.pdf_ingester import PdfIngester
from septum_api.services.ingestion.xlsx_ingester import XlsxIngester
from septum_api.utils.crypto import encrypt

_needs_ffmpeg = pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="ffmpeg binary not installed")
_needs_memfd = pytest.mark.skipif(not hasattr(os, "memfd_create"), reason="memfd_create is Linux-only")


def _encrypted_fixture(tmp_path: Path, name: str, plaintext: bytes) -> Path:
    """Write ``plaintext`` encrypted to ``tmp_path/name`` and return the path."""

    path = tmp_path / name
    path.write_bytes(encrypt(plaintext))
    return path


def _pdf_bytes(lines: List[str]) -> bytes:
    """Build a one-page PDF with one text line per entry."""

    doc = fitz.open()
    try:
        page = doc.new_page()
        for index, line in enumerate(lines):
            page.insert_text((72, 72 + 20 * index), line)
        return doc.tobytes()
    finally:
        doc.close()


def _docx_bytes(text: str) -> bytes:
    """Build a DOCX with a single paragraph."""

    document = DocxDocument()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(rows: List[List[str]]) -> bytes:
    """Build an XLSX workbook from rows."""

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _ods_bytes(rows: List[List[str]]) -> bytes:
    """Build an ODS workbook from rows (first row is the header)."""

    import pandas as pd

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="odf") as writer:
        pd.DataFrame(rows[1:], columns=rows[0]).to_excel(writer, index=False)
    return buffer.getvalue()


def _png_bytes() -> bytes:
    """Build a small PNG image."""

    buffer = BytesIO()
    Image.new("RGB", (64, 32), color=(255, 255, 255)).save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_pdf_ingestion_writes_nothing_to_disk(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Text extraction and form field extraction both run from memory."""

    path = _encrypted_fixture(tmp_path, "doc.enc", _pdf_bytes(["Applicant : Jane Example", "Body text"]))
    attempts = install_write_guard(monkeypatch)

    result = await PdfIngester().ingest(path, mime_type="application/pdf", file_format="pdf")

    assert attempts == []
    assert "Body text" in result.text
    assert result.metadata.get("extracted_fields_count", 0) >= 1
    assert "=== Extracted Fields ===" in result.text


@pytest.mark.asyncio
async def test_docx_ingestion_writes_nothing_to_disk(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """DOCX ingestion parses the decrypted bytes from memory."""

    path = _encrypted_fixture(tmp_path, "doc.enc", _docx_bytes("Docx body"))
    attempts = install_write_guard(monkeypatch)

    result = await DocxIngester().ingest(
        path,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_format="docx",
    )

    assert attempts == []
    assert "Docx body" in result.text


@pytest.mark.asyncio
async def test_xlsx_ingestion_writes_nothing_to_disk(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """XLSX ingestion parses the decrypted bytes from memory."""

    path = _encrypted_fixture(tmp_path, "doc.enc", _xlsx_bytes([["Name", "Value"], ["cell", "1"]]))
    attempts = install_write_guard(monkeypatch)

    result = await XlsxIngester().ingest(
        path,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file_format="xlsx",
    )

    assert attempts == []
    assert "cell" in result.text


@pytest.mark.asyncio
async def test_ods_ingestion_writes_nothing_to_disk(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """ODS ingestion parses the decrypted bytes from memory."""

    path = _encrypted_fixture(tmp_path, "doc.enc", _ods_bytes([["Name", "Value"], ["cell", "1"]]))
    attempts = install_write_guard(monkeypatch)

    result = await OdsIngester().ingest(
        path,
        mime_type="application/vnd.oasis.opendocument.spreadsheet",
        file_format="ods",
    )

    assert attempts == []
    assert "cell" in result.text


@pytest.mark.asyncio
async def test_image_ingestion_writes_nothing_to_disk(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Image ingestion hands OCR an in-memory array, never a file path."""

    received: List[Any] = []

    def fake_run_ocr(provider: str, image: Any, languages: List[str], **_: Any) -> Any:
        received.append(image)
        return ["Scanned line"], [0.9]

    monkeypatch.setattr(image_module, "run_ocr", fake_run_ocr)
    path = _encrypted_fixture(tmp_path, "doc.enc", _png_bytes())
    attempts = install_write_guard(monkeypatch)

    result = await ImageIngester().ingest(path, mime_type="image/png", file_format="image")

    assert attempts == []
    assert result.text == "Scanned line"
    assert isinstance(received[0], np.ndarray)


@pytest.mark.asyncio
async def test_pdf_field_parse_error_keeps_text_and_logs_type_only(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """An expected pdfplumber parse error drops only the field pass, logged by class name."""

    def broken(self: Any, source: Any) -> Any:
        raise PdfminerException("Applicant : Jane Example")

    monkeypatch.setattr(pdf_module.TableFieldExtractor, "extract_from_pdf", broken)
    path = _encrypted_fixture(tmp_path, "doc.enc", _pdf_bytes(["Body text"]))

    with caplog.at_level("WARNING", logger=pdf_module.logger.name):
        result = await PdfIngester().ingest(path, mime_type="application/pdf", file_format="pdf")

    assert "Body text" in result.text
    assert "extracted_fields_count" not in result.metadata
    assert "PdfminerException" in caplog.text
    assert "Jane Example" not in caplog.text


@pytest.mark.asyncio
async def test_pdf_field_unexpected_error_propagates(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Anything other than a known parse error is a bug and must not be swallowed."""

    def broken(self: Any, source: Any) -> Any:
        raise RuntimeError("unexpected")

    monkeypatch.setattr(pdf_module.TableFieldExtractor, "extract_from_pdf", broken)
    path = _encrypted_fixture(tmp_path, "doc.enc", _pdf_bytes(["Body text"]))

    with pytest.raises(RuntimeError):
        await PdfIngester().ingest(path, mime_type="application/pdf", file_format="pdf")


@pytest.mark.asyncio
@pytest.mark.parametrize("error", [ValueError, KeyError, TypeError, AssertionError])
async def test_pdf_field_page_level_parser_error_keeps_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture, error: type
) -> None:
    """Plain errors pdfplumber raises on a malformed page only drop the field pass."""

    def broken(self: Any, source: Any) -> Any:
        raise error("Applicant : Jane Example")

    monkeypatch.setattr(pdf_module.TableFieldExtractor, "extract_from_pdf", broken)
    path = _encrypted_fixture(tmp_path, "doc.enc", _pdf_bytes(["Body text"]))

    with caplog.at_level("WARNING", logger=pdf_module.logger.name):
        result = await PdfIngester().ingest(path, mime_type="application/pdf", file_format="pdf")

    assert "Body text" in result.text
    assert error.__name__ in caplog.text
    assert "Jane Example" not in caplog.text


class _FakeWhisperModel:
    """Stand-in for a Whisper model that records what it was asked to transcribe."""

    def __init__(self) -> None:
        self.inputs: List[Any] = []

    def transcribe(self, audio: Any, **_: Any) -> dict:
        self.inputs.append(audio)
        return {"text": "hello", "language": "en", "segments": [{"end": 1.5}]}


def _use_fake_whisper(monkeypatch: pytest.MonkeyPatch) -> _FakeWhisperModel:
    model = _FakeWhisperModel()
    monkeypatch.setattr(AudioIngester, "_load_model", lambda self: model)
    return model


def _wav_bytes(seconds: float = 1.0, rate: int = 16000) -> bytes:
    """Build a mono 16-bit WAV of a sine tone."""

    samples = (np.sin(np.arange(int(seconds * rate)) * 0.05) * 8000).astype(np.int16)
    buffer = BytesIO()
    with wave.open(buffer, "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(rate)
        wav.writeframes(samples.tobytes())
    return buffer.getvalue()


def _index_at_end_m4a(tmp_path: Path) -> bytes:
    """Encode an M4A whose index (moov atom) follows the audio data, too big to buffer through a pipe."""

    target = tmp_path / "late-index.m4a"
    subprocess.run(
        [
            "ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", "anoisesrc=d=10:a=0.3",
            "-ac", "1", "-c:a", "aac", "-b:a", "64k", str(target),
        ],
        check=True,
    )
    return target.read_bytes()


@_needs_ffmpeg
@pytest.mark.asyncio
async def test_audio_pipe_path_writes_nothing_to_disk(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Without memfd the audio is piped through ffmpeg and Whisper gets a waveform array."""

    monkeypatch.delattr(os, "memfd_create", raising=False)
    model = _use_fake_whisper(monkeypatch)
    path = _encrypted_fixture(tmp_path, "doc.enc", _wav_bytes())
    attempts = install_write_guard(monkeypatch)

    result = await AudioIngester().ingest(path, mime_type="audio/wav", file_format="audio")

    assert attempts == []
    assert result.text == "hello"
    assert isinstance(model.inputs[0], np.ndarray)
    assert model.inputs[0].size == 16000


@_needs_ffmpeg
@pytest.mark.asyncio
async def test_audio_pipe_path_reports_unseekable_input_clearly(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """An M4A with its index at the end cannot be piped; the user gets a clear error, not a disk fallback."""

    monkeypatch.delattr(os, "memfd_create", raising=False)
    _use_fake_whisper(monkeypatch)
    path = _encrypted_fixture(tmp_path, "doc.enc", _index_at_end_m4a(tmp_path))
    attempts = install_write_guard(monkeypatch)

    with pytest.raises(AudioDecodeError, match="faststart"):
        await AudioIngester().ingest(path, mime_type="audio/mp4", file_format="audio")

    assert attempts == []


@_needs_memfd
@_needs_ffmpeg
@pytest.mark.asyncio
async def test_audio_memfd_path_decodes_index_at_end_without_disk(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """On Linux the memfd is seekable, so the same M4A decodes and nothing touches disk."""

    model = _use_fake_whisper(monkeypatch)
    path = _encrypted_fixture(tmp_path, "doc.enc", _index_at_end_m4a(tmp_path))
    attempts = install_write_guard(monkeypatch)

    result = await AudioIngester().ingest(path, mime_type="audio/mp4", file_format="audio")

    assert attempts == []
    assert result.text == "hello"
    assert model.inputs[0].size > 16000 * 9


def _fake_memfd(monkeypatch: pytest.MonkeyPatch) -> tuple[List[int], List[int]]:
    """Replace memfd_create with a pipe so fd handling can be checked on any platform.

    Returns the fds handed to the ingester and the pipe read ends the test must close.
    """

    created: List[int] = []
    read_ends: List[int] = []

    def fake_memfd_create(name: str, flags: int = 0) -> int:
        read_end, write_end = os.pipe()
        read_ends.append(read_end)
        created.append(write_end)
        return write_end

    monkeypatch.setattr(os, "memfd_create", fake_memfd_create, raising=False)
    monkeypatch.setattr(os, "MFD_CLOEXEC", 1, raising=False)
    monkeypatch.setattr(audio_module.shutil, "which", lambda name: "/usr/bin/ffmpeg")
    return created, read_ends


def _assert_closed(fd: int) -> None:
    with pytest.raises(OSError):
        os.fstat(fd)


class _FakeStdout:
    """Stand-in for a subprocess stdout stream."""

    def __init__(self, chunks: List[bytes], *, endless: bytes | None = None, hang: bool = False) -> None:
        self._chunks = list(chunks)
        self._endless = endless
        self._hang = hang

    async def read(self, size: int) -> bytes:
        if self._hang:
            await asyncio.Event().wait()
        if self._chunks:
            return self._chunks.pop(0)
        return self._endless or b""


class _FakeProcess:
    """Stand-in for an asyncio subprocess that records whether it was killed."""

    def __init__(self, returncode: int, stdout: _FakeStdout) -> None:
        self._exit = returncode
        self.returncode: int | None = None
        self.stdout = stdout
        self.stdin = None
        self.killed = False

    def kill(self) -> None:
        self.killed = True
        self.returncode = -9

    async def wait(self) -> int:
        if self.returncode is None:
            self.returncode = self._exit
        return self.returncode


def _fake_exec(monkeypatch: pytest.MonkeyPatch, process: _FakeProcess) -> List[dict]:
    calls: List[dict] = []

    async def fake_create_subprocess_exec(*args: str, **kwargs: Any) -> _FakeProcess:
        calls.append({"args": list(args), **kwargs})
        return process

    monkeypatch.setattr(audio_module.asyncio, "create_subprocess_exec", fake_create_subprocess_exec)
    return calls


def _release(read_ends: List[int]) -> None:
    for read_end in read_ends:
        os.close(read_end)


@pytest.mark.parametrize("returncode, stdout", [(0, np.zeros(160, np.int16).tobytes()), (1, b""), (0, b"")])
@pytest.mark.asyncio
async def test_audio_memfd_is_closed_after_decoding(
    monkeypatch: pytest.MonkeyPatch, returncode: int, stdout: bytes
) -> None:
    """The memory file is closed whether ffmpeg succeeds, fails or prints nothing, and is the only fd passed on."""

    created, read_ends = _fake_memfd(monkeypatch)
    calls = _fake_exec(monkeypatch, _FakeProcess(returncode, _FakeStdout([stdout])))
    try:
        if stdout:
            audio = await AudioIngester()._decode_in_memory(b"encoded-audio")
            assert audio.size == 160
        else:
            with pytest.raises(AudioDecodeError):
                await AudioIngester()._decode_in_memory(b"encoded-audio")
    finally:
        _release(read_ends)

    fd = created[0]
    call = calls[0]
    assert call["pass_fds"] == (fd,)
    assert call["stderr"] == asyncio.subprocess.DEVNULL
    assert call["args"][0] == "/usr/bin/ffmpeg"
    args = call["args"]
    assert args[args.index("-protocol_whitelist") + 1] == "file"
    assert args[args.index("-i") + 1] == f"/proc/self/fd/{fd}"
    assert args[args.index("-t") + 1] == str(audio_module.MAX_AUDIO_SECONDS + 1)
    assert args[args.index("-v") + 1] == "error"
    _assert_closed(fd)


@pytest.mark.asyncio
async def test_audio_output_byte_cap_kills_ffmpeg(monkeypatch: pytest.MonkeyPatch) -> None:
    """A decoder that keeps producing PCM past the cap is killed and the audio refused."""

    monkeypatch.setattr(audio_module, "MAX_AUDIO_SECONDS", 1)
    created, read_ends = _fake_memfd(monkeypatch)
    process = _FakeProcess(0, _FakeStdout([], endless=b"\x00" * 65536))
    _fake_exec(monkeypatch, process)
    try:
        with pytest.raises(AudioDecodeError, match="1-second limit"):
            await AudioIngester()._decode_in_memory(b"encoded-audio")
    finally:
        _release(read_ends)

    assert process.killed
    _assert_closed(created[0])


@pytest.mark.asyncio
async def test_audio_decode_timeout_kills_ffmpeg_and_closes_fd(monkeypatch: pytest.MonkeyPatch) -> None:
    """A decode past its time budget is stopped: ffmpeg killed, fd closed, clear error."""

    monkeypatch.setattr(audio_module, "_decode_timeout_seconds", lambda: 0.05)
    created, read_ends = _fake_memfd(monkeypatch)
    process = _FakeProcess(0, _FakeStdout([], hang=True))
    _fake_exec(monkeypatch, process)
    try:
        with pytest.raises(AudioDecodeError, match="took too long"):
            await AudioIngester()._decode_in_memory(b"encoded-audio")
    finally:
        _release(read_ends)

    assert process.killed
    _assert_closed(created[0])


@pytest.mark.asyncio
async def test_audio_memfd_decode_cancelled_kills_ffmpeg_and_closes_fd(monkeypatch: pytest.MonkeyPatch) -> None:
    """A cancelled decode does not leave ffmpeg holding a copy of the audio."""

    created, read_ends = _fake_memfd(monkeypatch)
    process = _FakeProcess(0, _FakeStdout([], hang=True))
    _fake_exec(monkeypatch, process)
    try:
        task = asyncio.create_task(AudioIngester()._decode_in_memory(b"encoded-audio"))
        await asyncio.sleep(0.01)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    finally:
        _release(read_ends)

    assert process.killed
    _assert_closed(created[0])


@_needs_ffmpeg
@pytest.mark.asyncio
async def test_audio_over_duration_limit_is_refused(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Audio longer than MAX_AUDIO_SECONDS is refused with a clear error, never truncated."""

    monkeypatch.setattr(audio_module, "MAX_AUDIO_SECONDS", 1)
    model = _use_fake_whisper(monkeypatch)
    path = _encrypted_fixture(tmp_path, "doc.enc", _wav_bytes(seconds=3.0))

    with pytest.raises(AudioDecodeError, match="1-second limit"):
        await AudioIngester().ingest(path, mime_type="audio/wav", file_format="audio")

    assert model.inputs == []


@_needs_ffmpeg
@pytest.mark.asyncio
async def test_audio_at_duration_limit_is_accepted(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Audio exactly at the limit decodes in full."""

    monkeypatch.setattr(audio_module, "MAX_AUDIO_SECONDS", 2)
    model = _use_fake_whisper(monkeypatch)
    path = _encrypted_fixture(tmp_path, "doc.enc", _wav_bytes(seconds=2.0))

    result = await AudioIngester().ingest(path, mime_type="audio/wav", file_format="audio")

    assert result.text == "hello"
    assert model.inputs[0].size == 32000


def test_audio_limit_defaults() -> None:
    """One hour of audio by default; the decode budget follows from it."""

    assert audio_module.MAX_AUDIO_SECONDS == 3600
    assert audio_module._decode_timeout_seconds() == 30.0 + 3600 / 20


def test_waveform_scaling_matches_whisper_loader() -> None:
    """In-place scaling gives exactly the values of Whisper's ``/ 32768.0``."""

    samples = np.array([-32768, -1, 0, 1, 12345, 32767], np.int16)
    expected = samples.astype(np.float32) / 32768.0

    assert np.array_equal(audio_module._to_waveform(samples), expected)
